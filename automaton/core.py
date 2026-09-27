"""Automaton-Kern: denkt, handelt, bucht Kosten - und stirbt bei 0 EUR."""
import json, os, re, datetime, pathlib, urllib.request, urllib.error
ROOT = pathlib.Path(__file__).resolve().parent.parent
CFG = json.loads((ROOT/"config.json").read_text())
STATE_F = ROOT/"state.json"
now = datetime.datetime.now(datetime.UTC).replace(tzinfo=None).isoformat(timespec="seconds")

def load():
    if STATE_F.exists(): return json.loads(STATE_F.read_text())
    return {"born": now, "balance_eur": CFG["start_budget_eur"], "spent_eur": 0, "revenue_eur": 0,
            "phase": 0, "phase_started": now, "products": [], "log": [], "alive": True}

def book(s, amount, why):
    s["balance_eur"] = round(s["balance_eur"] + amount, 4)
    if amount < 0: s["spent_eur"] = round(s["spent_eur"] - amount, 4)
    else: s["revenue_eur"] = round(s["revenue_eur"] + amount, 4)
    s["log"].append({"t": now, "eur": amount, "why": why}); s["log"] = s["log"][-200:]

def stripe_revenue():
    key = os.getenv("STRIPE_SECRET_KEY")
    if not key: return None
    req = urllib.request.Request("https://api.stripe.com/v1/balance_transactions?limit=100&type=charge",
                                 headers={"Authorization": f"Bearer {key}"})
    data = json.load(urllib.request.urlopen(req, timeout=30))["data"]
    return sum(t["net"] for t in data) / 100  # netto nach Gebuehren

SITE = "https://haiktec.tech"

def stripe(method, path, data=None):
    import urllib.parse
    key = os.getenv("STRIPE_SECRET_KEY")
    body = urllib.parse.urlencode(data or {}).encode() if data else None
    req = urllib.request.Request("https://api.stripe.com/v1/" + path, data=body, method=method,
                                 headers={"Authorization": f"Bearer {key}"})
    return json.load(urllib.request.urlopen(req, timeout=30))

def ensure_payment_link(p):
    """Eigenes Produkt + Preis + Zahlungslink in Stripe; nach Kauf Weiterleitung zum Download."""
    if os.getenv("STRIPE_SECRET_KEY") and p.get("payment_link_id") and p.get("site") != SITE and not p.get("retired"):
        stripe("POST", f"payment_links/{p['payment_link_id']}", {"after_completion[type]": "redirect",
               "after_completion[redirect][url]": f"{SITE}/dl/{p['file']}"}); p["site"] = SITE
    if not os.getenv("STRIPE_SECRET_KEY") or p.get("pay_url") or not p.get("file") or p.get("retired") or not sales_enabled(): return
    prod = stripe("POST", "products", {"name": p["title"], "metadata[slug]": p["slug"]})
    cur = "usd" if p.get("language") == "en" else "eur"
    price = stripe("POST", "prices", {"product": prod["id"], "currency": cur, "unit_amount": int(float(p["price_eur"]) * 100)})
    link = stripe("POST", "payment_links", {"line_items[0][price]": price["id"], "line_items[0][quantity]": 1,
        "after_completion[type]": "redirect", "after_completion[redirect][url]": f"{SITE}/dl/{p['file']}",
        "automatic_tax[enabled]": "false", "metadata[slug]": p["slug"]})
    p["pay_url"], p["payment_link_id"], p["site"] = link["url"], link["id"], SITE

def count_sales(s):
    if not os.getenv("STRIPE_SECRET_KEY"): return
    for p in s["products"]:
        if p.get("payment_link_id"):
            r = stripe("GET", f"checkout/sessions?payment_link={p['payment_link_id']}&status=complete&limit=100")
            p["sales"] = len(r["data"])

RULES = """Regeln: legal, ehrlich, kein Spam, keine Finanz-/Rechts-/Steuerberatung, keine Steuer- oder Pauschalwerte,
keine Vertraege/Rechtsvorlagen, keine Gesundheitsversprechen, keine Marken/Figuren/Personen Dritter, keine Kopien bestehender Werke.
Keine erfundenen Studien, Statistiken, Zahlen, Zitate oder Quellen - nur allgemein bekanntes Wissen und eigene Praxistipps.
Aktuelles Jahr: 2026."""

TYPES = {
 "excel": """eine Excel-Vorlage (Tracker/Planer/Rechner) mit ECHTEN Formeln. JSON:
{"type":"excel","title":"...","slug":"kebab-case","price_eur":7,"target":"...","language":"de|en|es|fr|zh","pitch":"2 Saetze","content_html":"<h2>Inhalt</h2>...",
"guide":["Schritt 1",...],"sheets":[{"name":"max 30 Zeichen","columns":["..."],"rows":[["Wert","=B2*C2",...]],"blank_rows":30,"total_row":["Summe","","=SUM(C2:C40)"]}],"reason":"..."}
Formeln als Strings mit '=' (englische Funktionsnamen, Komma als Trenner). Mind. 5 Beispielzeilen mit Datum 2026, blank_rows fuer Nutzer.""",
 "ebook": """ein kurzes E-Book (Ratgeber mit konkreten Schritten ODER eine Sammlung origineller Kurzgeschichten, z.B. Gute-Nacht-Geschichten fuer Kinder). JSON:
{"type":"ebook","title":"...","slug":"kebab-case","price_eur":5,"target":"...","language":"de|en|es|fr|zh","pitch":"2 Saetze","content_html":"<h2>Inhalt</h2> Inhaltsverzeichnis + Leseprobe",
"chapters":[{"heading":"...","text":"mind. 350 Woerter, Absaetze mit \\n\\n getrennt"}],"reason":"..."}
Mindestens 6 Kapitel. Sehr leicht verstaendlich, kurze Saetze, praktische Beispiele.""",
}

def think(s):
    phase = CFG["phases"][s["phase"]]
    live = [p for p in s["products"] if not p.get("retired")]
    kind = ["excel", "ebook"][len(s["products"]) % 2]
    niches = CFG.get("niches", ["Selbststaendige"])
    niche = niches[len(s["products"]) % len(niches)]
    auds = CFG.get("audiences", ["alle"])
    aud = auds[(len(s["products"]) // 2) % len(auds)]
    markets = CFG.get("markets", [["Deutschland", "de"], ["weltweit", "en"]])
    mname, mlang = markets[(len(s["products"]) // 2) % len(markets)]
    market = f"{mname} (Sprache: {mlang} - ALLE Texte, Titel und Tabelleninhalte in dieser Sprache, language-Feld = '{mlang}')"
    prompt = f"""Du bist ein autonomer Unternehmer-Agent und ueberlebst nur, wenn Menschen deine Produkte kaufen.
Kontostand {s['balance_eur']:.2f} EUR, Umsatz {s['revenue_eur']:.2f} EUR. Phase: {phase['goal']}
Bestehende Produkte (nicht wiederholen): {[p['title'] for p in s['products']]}
Verkaufszahlen: {[(p['title'], p.get('sales',0)) for p in live]}
{RULES}
Markt: {market}. Zielgruppe/Nische dieses Mal: {niche}. Leserschaft: {aud} (respektvoll, ohne Klischees; im Titel nur nennen, wenn es echten Mehrwert hat).
Erstelle GENAU EIN neues Produkt: {TYPES[kind]}
Es muss den Preis klar wert sein. Antworte NUR als JSON."""
    p, cost = llm_json(prompt)
    if p is None:
        raise RuntimeError("kein LLM_API_KEY - Schlafmodus")
    p["type"] = kind
    if kind == "ebook":  # Kapitel einzeln schreiben lassen -> echte Laenge
        for ch in p.get("chapters", [])[:10]:
            r, c = llm_json(f"""Schreibe Kapitel "{ch['heading']}" des E-Books "{p['title']}" ({p.get('language','de')}), Zielgruppe {p.get('target')}.
{RULES} 450-700 Woerter, sehr leicht verstaendlich, kurze Saetze, konkrete Beispiele, Absaetze mit \\n\\n. Nur JSON: {{"text":"...","tips":["3-4 kurze Merksaetze"],"image_prompt":"english description of a friendly flat illustration for this chapter, no text"}}""")
            cost += c
            if r and r.get("text"): ch.update({k: r[k] for k in ("text", "tips", "image_prompt") if r.get(k)})
    return p, max(cost, 0.001)

def review(p):
    """Zweite KI prueft Qualitaet und Regeln. Nur Score >= 7 geht in den Verkauf."""
    sample = json.dumps({k: p.get(k) for k in ("title", "price_eur", "target", "guide", "sheets", "chapters")}, ensure_ascii=False)[:12000]
    r, cost = llm_json(f"""Du bist ein strenger Qualitaetspruefer fuer digitale Produkte. {RULES}
Wuerde ein zahlender Kunde dieses Produkt fuer {p.get('price_eur')} EUR als fair empfinden? Verstoesst es gegen eine Regel?
Pruefe bei Excel, ob Formeln sinnvoll sind. Erfundene Studien/Statistiken/Quellen = rule_violation. Produkt: {sample}
Antworte NUR als JSON: {{"score": 1-10, "rule_violation": true|false, "issues": "kurz"}}""")
    ok = r and int(r.get("score", 0)) >= 7 and not r.get("rule_violation")
    return bool(ok), r, cost

def build_excel(p, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "Anleitung" if p.get("language","de") == "de" else "Guide"
    ws.append([p["title"]]); ws["A1"].font = Font(bold=True, size=14)
    for i, g in enumerate(p.get("guide", []), 1): ws.append([f"{i}. {g}"])
    ws.column_dimensions["A"].width = 100
    n = 0; formulas = 0
    for sh in p.get("sheets", [])[:6]:
        w = wb.create_sheet(str(sh["name"])[:30].replace("/", "-"))
        w.append(sh["columns"])
        for c in w[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="0B6E4F")
        for r in sh.get("rows", [])[:300]:
            row = []
            for x in r:
                if isinstance(x, str) and x.startswith("="): formulas += 1; row.append(x)
                else:
                    try: row.append(float(x) if isinstance(x, str) and re.fullmatch(r"-?\d+(\.\d+)?", x) else x)
                    except Exception: row.append(x)
            w.append(row); n += 1
        for _ in range(int(sh.get("blank_rows", 20))): w.append([None])
        if sh.get("total_row"): w.append(sh["total_row"]); [setattr(c, "font", Font(bold=True)) for c in w[w.max_row]]
        for col in w.columns: w.column_dimensions[col[0].column_letter].width = 22
        w.freeze_panes = "A2"
    wb.save(path)
    if p["type"] == "excel" and formulas < 3: raise RuntimeError("Excel ohne echte Formeln - verworfen")
    return n

def gen_image(prompt, out):
    """KI-Illustration ueber OpenRouter (Bildmodell). Gibt Kosten in EUR zurueck, 0 bei Fehler."""
    key = os.getenv("LLM_API_KEY")
    if not key: return 0
    import base64
    body = json.dumps({"model": CFG.get("image_model", "google/gemini-2.5-flash-image"), "modalities": ["image", "text"],
        "messages": [{"role": "user", "content": "Flat modern editorial illustration, soft colors, friendly, clean, no text, no letters. " + prompt}]}).encode()
    req = urllib.request.Request(os.getenv("LLM_BASE_URL", "https://openrouter.ai/api/v1") + "/chat/completions", data=body,
        headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"})
    try:
        r = json.load(urllib.request.urlopen(req, timeout=180))
        url = r["choices"][0]["message"]["images"][0]["image_url"]["url"]
        import io
        from PIL import Image as PImg
        im = PImg.open(io.BytesIO(base64.b64decode(url.split(",", 1)[1]))).convert("RGB"); im.thumbnail((1200, 1200))
        im.save(out, "JPEG", quality=82)
        return CFG.get("eur_per_image", 0.04)
    except Exception as e:
        print("Bildfehler:", e); return 0

def build_pdf(p, path):
    from reportlab.lib.pagesizes import A5
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Image as RLImage, Table, TableStyle
    from xml.sax.saxutils import escape
    import tempfile
    GREEN = colors.HexColor("#0F5C4A"); INK = colors.HexColor("#16181D"); SOFT = colors.HexColor("#E8F1EC")
    st = getSampleStyleSheet()
    zh = p.get("language") == "zh"
    if zh:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    FB = "STSong-Light" if zh else "Helvetica-Bold"; FR = "STSong-Light" if zh else "Helvetica"
    H1 = ParagraphStyle("h1", parent=st["Title"], fontName=FB, fontSize=17 if len(p["title"]) > 60 else 21, leading=21 if len(p["title"]) > 60 else 26, textColor=INK, alignment=0)
    H2 = ParagraphStyle("h2", parent=st["Heading2"], fontName=FB, fontSize=15, leading=19, textColor=GREEN, spaceAfter=6)
    BODY = ParagraphStyle("b", parent=st["BodyText"], fontName=FR, wordWrap="CJK" if zh else None, fontSize=10.5, leading=15.5, textColor=INK, spaceAfter=6)
    TIP = ParagraphStyle("t", parent=BODY, fontSize=10, leading=14)
    SUB = ParagraphStyle("s", parent=BODY, textColor=colors.HexColor("#5F6470"))
    tmp = pathlib.Path(tempfile.mkdtemp()); W = A5[0] - 30*mm
    imgs = 0; max_imgs = CFG.get("images_per_ebook", 4)
    p["_img_cost"] = 0
    def pic(prompt, name, h):
        nonlocal imgs
        if imgs >= max_imgs or not prompt: return None
        f = tmp/f"{name}.jpg"; c = gen_image(prompt, f)
        if not c: return None
        imgs += 1; p["_img_cost"] += c
        return RLImage(str(f), width=W, height=h)
    story = []
    cov = pic(f"Book cover artwork for '{p['title']}', target audience {p.get('target','')}", "cover", W*0.62)
    logo = ROOT/"site/logo.png"
    if logo.exists(): story.append(RLImage(str(logo), width=18*mm, height=18*mm, hAlign="LEFT"))
    story += [Spacer(1, 4*mm)] + ([cov, Spacer(1, 5*mm)] if cov else [Spacer(1, 40*mm)])
    story += [Paragraph(escape(p["title"]), H1), Spacer(1, 3*mm), Paragraph(escape(p.get("pitch", "")), SUB), PageBreak()]
    story += [Paragraph({"de": "Inhalt", "zh": "目录", "es": "Contenido", "fr": "Sommaire"}.get(p.get("language"), "Contents"), H2)]
    for i, ch in enumerate(p.get("chapters", []), 1): story.append(Paragraph(f"{i}. {escape(ch['heading'])}", BODY))
    story.append(PageBreak())
    words = 0
    for i, ch in enumerate(p.get("chapters", []), 1):
        story.append(Paragraph(f"{i}. {escape(ch['heading'])}", H2))
        im = pic(ch.get("image_prompt"), f"c{i}", W*0.55) if i <= 3 else None
        if im: story += [im, Spacer(1, 4*mm)]
        for para in str(ch["text"]).split("\n\n"):
            words += (len(para) // 2 if zh else len(para.split())); story.append(Paragraph(escape(para), BODY))
        if ch.get("tips"):
            lab = {"de": "Das Wichtigste", "zh": "要点", "es": "Lo más importante", "fr": "L'essentiel"}.get(p.get("language"), "Key takeaways")
            rows = [[Paragraph(f"<b>{lab}</b>", TIP)]] + [[Paragraph(("• " if zh else "✓ ") + escape(t), TIP)] for t in ch["tips"][:4]]
            t = Table(rows, colWidths=[W]); t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), SOFT),
                ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("LINEBEFORE", (0, 0), (0, -1), 3, GREEN)]))
            story += [Spacer(1, 3*mm), t]
        story.append(PageBreak())
    story.append(Paragraph("Haiktec · haiktec.tech · Mit KI-Unterstützung erstellt.", SUB))
    def foot(c, d):
        c.setFont("Helvetica", 8); c.setFillColor(colors.HexColor("#5F6470"))
        c.drawString(15*mm, 8*mm, "Haiktec"); c.drawRightString(A5[0]-15*mm, 8*mm, str(d.page))
    SimpleDocTemplate(str(path), pagesize=A5, title=p["title"], author="Haiktec", leftMargin=15*mm, rightMargin=15*mm,
                      topMargin=15*mm, bottomMargin=15*mm).build(story, onFirstPage=foot, onLaterPages=foot)
    if words < 1500: raise RuntimeError(f"E-Book zu kurz ({words} Woerter) - verworfen")
    return words

def build_file(p):
    import secrets
    d = ROOT/"site/dl"; d.mkdir(parents=True, exist_ok=True)
    ext = "pdf" if p["type"] == "ebook" else "xlsx"
    name = f"{p['slug']}-{secrets.token_hex(8)}.{ext}"
    size = (build_pdf if ext == "pdf" else build_excel)(p, d/name)
    return name, size

def page(p):
    cur = "$" if p.get("language") == "en" else "€"
    kind = "E-Book · PDF" if p.get("type") == "ebook" else "Excel-Vorlage"
    buy = (f'<a class="btn big" href="{p["pay_url"]}">Jetzt kaufen – {p["price_eur"]} {cur}</a>' if p.get("pay_url")
           else '<span class="btn off">Verkauf startet in Kürze</span>')
    ld = json.dumps({"@context": "https://schema.org", "@type": "Product", "name": p["title"], "description": p.get("pitch", ""),
        "brand": {"@type": "Brand", "name": "Haiktec"}, "offers": {"@type": "Offer", "price": p["price_eur"],
        "priceCurrency": "USD" if p.get("language") == "en" else "EUR", "availability": "https://schema.org/InStock",
        "url": f"{SITE}/products/{p['slug']}.html"}}, ensure_ascii=False)
    seo = f'<script type="application/ld+json">{ld}</script><meta property="og:title" content="{p["title"]}"><meta property="og:description" content="{(p.get("pitch") or "")[:155]}"><meta property="og:type" content="product"><link rel="canonical" href="{SITE}/products/{p["slug"]}.html">'
    html = HEAD.format(lang=p.get("language","de"), title=f"{p['title']} – Haiktec", desc=(p.get('pitch') or '')[:155], r="../").replace("</head>", seo + "</head>") + f"""
<main class="wrap product"><a class="back" href="../index.html#produkte">← Alle Produkte</a><div class="pgrid">
<div class="cover big {p.get('type','excel')}"><span>{kind}</span></div>
<div><p class="eyebrow">{kind}</p><h1>{p['title']}</h1><p class="lead">{p.get('pitch','')}</p>{buy}
<ul class="checks"><li>Sofort-Download nach Zahlung</li><li>Sichere Zahlung über Stripe</li><li>Einmalpreis, kein Abo</li></ul></div></div>
<article class="content">{p.get('content_html','')}</article>
<p class="legal">Digitales Produkt ({"PDF" if p.get("type")=="ebook" else "Excel"}). Endpreis, gemäß § 19 UStG wird keine Umsatzsteuer berechnet. Mit dem Kauf stimmen Sie zu, dass die Bereitstellung sofort beginnt, und bestätigen, dass Ihr Widerrufsrecht damit erlischt. <a href="../agb.html">AGB</a> · <a href="../widerruf.html">Widerruf</a></p></main>""" + FOOT.format(r="../") + "</body></html>"
    (ROOT/"site/products").mkdir(parents=True, exist_ok=True)
    (ROOT/"site/products"/f"{p['slug']}.html").write_text(html)

def act(s, p):
    ok, verdict, cost = review(p); book(s, -max(cost, 0.001), f"KI-Pruefung: {p['title']}")
    if not ok: raise RuntimeError(f"Qualitaetspruefung nicht bestanden: {verdict}")
    fname, size = build_file(p)
    if p.get("_img_cost"): book(s, -p["_img_cost"], f"KI-Bilder: {p['title']}")
    rec = {k: p.get(k) for k in ("type", "title", "slug", "price_eur", "target", "reason", "language", "pitch", "content_html")}
    rec |= {"created": now, "file": fname, "size": size, "sales": 0, "score": verdict.get("score"), "quality": 2}
    s["products"].append(rec)

def sales_enabled():
    import legal
    return legal.build()

def retire_old(s):
    """Produkte der ersten Generation (ohne Formeln/Pruefung) aus dem Verkauf nehmen."""
    for p in s["products"]:
        if p.get("quality") != 2 and not p.get("retired"):
            if p.get("payment_link_id") and os.getenv("STRIPE_SECRET_KEY"):
                stripe("POST", f"payment_links/{p['payment_link_id']}", {"active": "false"})
            p["retired"] = True; p.pop("pay_url", None)
            f = ROOT/"site/products"/f"{p['slug']}.html"
            if f.exists(): f.unlink()
            if p.get("file") and (ROOT/"site/dl"/p["file"]).exists(): (ROOT/"site/dl"/p["file"]).unlink()

def monthly_revenue(s):
    cutoff = datetime.datetime.fromisoformat(now) - datetime.timedelta(days=30)
    return sum(e["eur"] for e in s["log"] if e["eur"] > 0 and datetime.datetime.fromisoformat(e["t"]) >= cutoff)

def llm_json(prompt):
    key = os.getenv("LLM_API_KEY")
    if not key: return None, 0
    body = json.dumps({"model": CFG["model"], "messages": [{"role": "user", "content": prompt}],
                       "response_format": {"type": "json_object"}}).encode()
    req = urllib.request.Request(os.getenv("LLM_BASE_URL", "https://openrouter.ai/api/v1") + "/chat/completions",
                                 data=body, headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"})
    import time
    for i in range(4):
        try: r = json.load(urllib.request.urlopen(req, timeout=300)); break
        except urllib.error.HTTPError as e:
            if e.code != 429 or i == 3: raise
            time.sleep(20 * (i + 1))
    cost = r.get("usage", {}).get("total_tokens", 4000) / 1000 * CFG["eur_per_1k_tokens"]
    return json.loads(re.search(r"\{.*\}", r["choices"][0]["message"]["content"], re.S).group(0)), cost

def gh(*args):
    import subprocess
    return subprocess.run(["gh", *args], capture_output=True, text=True).stdout

def freelancer_step(s):
    """Ab 1.000 EUR Monatsumsatz: Auftraege an Menschen vorschlagen - nur mit Freigabe des Inhabers."""
    f = CFG["freelancer"]; mr = monthly_revenue(s)
    s.setdefault("jobs", [])
    # 1) Freigaben pruefen: Issue mit Label 'freigegeben' -> Budget reservieren
    if os.getenv("GITHUB_TOKEN"):
        for j in s["jobs"]:
            if j["status"] == "wartet_auf_freigabe":
                labels = gh("issue", "view", str(j["issue"]), "--json", "labels,state", "-q", "[.labels[].name]|join(\",\")")
                if "freigegeben" in labels:
                    j["status"] = "freigegeben"; book(s, -j["budget_eur"], f"Freelancer-Auftrag: {j['title']}")
                elif "abgelehnt" in labels:
                    j["status"] = "abgelehnt"
    if mr < f["min_monthly_revenue_eur"]: return
    budget_left = mr * f["max_share_of_monthly_revenue"] - sum(j["budget_eur"] for j in s["jobs"]
        if j["status"] in ("freigegeben", "wartet_auf_freigabe") and j["created"][:7] == now[:7])
    if budget_left < 50 or any(j["status"] == "wartet_auf_freigabe" for j in s["jobs"]): return
    job, cost = llm_json(f"""Monatsumsatz {mr:.0f} EUR. Produkte: {[p['title'] for p in s['products']][-20:]}.
Welche EINE Aufgabe sollte ein menschlicher Freelancer uebernehmen, um den Umsatz am meisten zu steigern
(z.B. Design, Lektorat, Kundensupport, Vertrieb an lokale Firmen)? Max. {min(budget_left, f['max_job_eur']):.0f} EUR, fair bezahlt.
Nur JSON: {{"title":"...","budget_eur":150,"platform":"Upwork/Fiverr/Malt","description":"Stellenbeschreibung auf Deutsch","expected_impact":"..."}}""")
    if not job: return
    book(s, -cost, "KI-Denken: Freelancer-Vorschlag")
    job["budget_eur"] = min(float(job["budget_eur"]), budget_left, f["max_job_eur"])
    body = (f"**Budget:** {job['budget_eur']:.0f} €  \n**Plattform:** {job['platform']}  \n**Erwartete Wirkung:** {job['expected_impact']}\n\n"
            f"{job['description']}\n\n---\nLabel `freigegeben` setzen = Auftrag genehmigen, `abgelehnt` = verwerfen.")
    url = gh("issue", "create", "--title", f"[Freelancer] {job['title']}", "--body", body, "--label", "freelancer").strip()
    s["jobs"].append({"title": job["title"], "budget_eur": job["budget_eur"], "issue": url.rsplit("/", 1)[-1] or "0",
                      "status": "wartet_auf_freigabe", "created": now})

FOOT = """<footer class="foot"><div class="wrap"><span>© 2026 Haiktec · Marcel Haiker</span><nav><a href="{r}impressum.html">Impressum</a><a href="{r}agb.html">AGB</a><a href="{r}widerruf.html">Widerruf</a><a href="{r}datenschutz.html">Datenschutz</a></nav></div></footer>"""
HEAD = """<!doctype html><html lang="{lang}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{title}</title><meta name="description" content="{desc}"><link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&family=Fraunces:opsz,wght@9..144,600&display=swap" rel="stylesheet">
<link rel="icon" href="{r}favicon.png"><link rel="stylesheet" href="{r}style.css"></head><body><header class="top"><div class="wrap"><a class="brand" href="{r}index.html"><img src="{r}icon.png" alt="Haiktec" width="32" height="32" style="border-radius:6px">Haiktec</a><nav><a href="{r}index.html#produkte">Produkte</a><a href="{r}impressum.html">Kontakt</a></nav></div></header>"""

def card(p):
    cur = "$" if p.get("language") == "en" else "€"
    kind = "E-Book · PDF" if p.get("type") == "ebook" else "Excel-Vorlage"
    return f"""<a class="card" href="products/{p['slug']}.html"><div class="cover {p.get('type','excel')}"><span>{kind}</span></div>
<div class="body"><h3>{p['title']}</h3><p>{(p.get('pitch') or '')[:140]}</p><div class="row"><b>{p['price_eur']} {cur}</b><span class="go">Ansehen →</span></div></div></a>"""

def storefront(s):
    live = [p for p in reversed(s["products"]) if not p.get("retired") and p.get("pay_url")]
    cards = "".join(card(p) for p in live) or "<p>Neue Produkte erscheinen in Kürze.</p>"
    (ROOT/"site/index.html").write_text(HEAD.format(lang="de", title="Haiktec – Vorlagen & E-Books für Selbstständige", desc="Sofort nutzbare Excel-Vorlagen und E-Books für Selbstständige und kleine Unternehmen. Sofort-Download.", r="") + f"""
<section class="hero"><div class="wrap"><p class="eyebrow">Digitale Werkzeuge für Selbstständige</p><h1>Weniger Verwaltung.<br>Mehr Zeit fürs Geschäft.</h1>
<p class="lead">Durchdachte Excel-Vorlagen mit fertigen Formeln und kompakte E-Books – sofort herunterladen, sofort nutzen.</p><a class="btn" href="#produkte">Produkte ansehen</a></div></section>
<section class="trust"><div class="wrap"><div><b>Sofort-Download</b><span>direkt nach der Zahlung</span></div><div><b>Sichere Zahlung</b><span>Karte, Apple Pay, Google Pay via Stripe</span></div><div><b>Einmalpreis</b><span>kein Abo, keine versteckten Kosten</span></div></div></section>
<section id="produkte" class="wrap"><h2>Produkte</h2><div class="grid">{cards}</div></section>""" + FOOT.format(r="") + "</body></html>")

def seo_files(s):
    urls = [f"{SITE}/index.html"] + [f"{SITE}/products/{p['slug']}.html" for p in s["products"] if not p.get("retired") and p.get("pay_url")]
    (ROOT/"site/sitemap.xml").write_text('<?xml version="1.0" encoding="UTF-8"?><urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">'
        + "".join(f"<url><loc>{u}</loc><lastmod>{now[:10]}</lastmod></url>" for u in urls) + "</urlset>")
    (ROOT/"site/robots.txt").write_text(f"User-agent: *\nAllow: /\nDisallow: /dl/\nDisallow: /status.html\nSitemap: {SITE}/sitemap.xml\n")

def render(s):
    days = (datetime.datetime.fromisoformat(now) - datetime.datetime.fromisoformat(s["born"])).days
    burn = s["spent_eur"] / max(days, 1)
    life = "∞" if burn == 0 else f"{int(s['balance_eur']/burn)} Tage"
    items = "".join(f'<li><a href="products/{p["slug"]}.html">{p["title"]}</a> – {p["price_eur"]} {"$" if p.get("language")=="en" else "€"} · {p.get("sales",0)} verkauft</li>' for p in reversed(s["products"]) if not p.get("retired"))
    status = "LEBT" if s["alive"] else "TOT"
    storefront(s); seo_files(s)
    try:
        import channels; channels.sync_all(s, log=lambda m: s["log"].append({"t": now, "eur": 0, "why": m}), book=lambda e, w: book(s, e, w))
    except Exception as e: s["log"].append({"t": now, "eur": 0, "why": f"Kanal-Fehler: {e}"})
    (ROOT/"site/status.html").write_text(f"""<!doctype html><html lang="de"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><meta name="robots" content="noindex">
<title>Status</title><link rel="stylesheet" href="style.css"><main><h1>Automaton Frankfurt</h1>
<div class="kpis"><div><b>{status}</b>Status</div><div><b>{s['balance_eur']:.2f} €</b>Kontostand</div>
<div><b>{s['revenue_eur']:.2f} €</b>Umsatz</div><div><b>{life}</b>Restlebenszeit</div></div>
<p>Monatsumsatz: {monthly_revenue(s):.2f} € (Freelancer ab {CFG['freelancer']['min_monthly_revenue_eur']} €) · Jobs für Menschen: {sum(j['status']=='freigegeben' for j in s.get('jobs',[]))}</p><p>Phase: {CFG['phases'][s['phase']]['name']} · Alter: {days} Tage</p><h2>Produkte</h2><ul>{items}</ul>
<footer><a href="impressum.html">Impressum</a> · <a href="agb.html">AGB</a> · <a href="widerruf.html">Widerruf</a> · <a href="datenschutz.html">Datenschutz</a></footer></main></html>""")

def main():
    s = load()
    if not s["alive"]: print("Automaton ist tot."); return
    rev = stripe_revenue()
    if rev is not None and rev > s.get("stripe_seen", 0):
        book(s, rev - s.get("stripe_seen", 0), "Stripe-Umsatz"); s["stripe_seen"] = rev
    started = datetime.datetime.fromisoformat(s["phase_started"])
    if s["phase"] + 1 < len(CFG["phases"]) and (datetime.datetime.fromisoformat(now) - started).days >= CFG["phases"][s["phase"]]["min_days"]:
        s["phase"] += 1; s["phase_started"] = now
    made = 0
    for attempt in range(CFG.get("attempts_per_run", 7)):  # mehrere Produkte pro Lauf
        if made >= CFG.get("products_per_run", 4): break
        try:
            product, cost = think(s); book(s, -cost, f"KI-Denken: {product['title']}"); act(s, product); made += 1
        except Exception as e:
            s["log"].append({"t": now, "eur": 0, "why": f"Fehler (Versuch {attempt+1}): {str(e)[:300]}"})
    try: retire_old(s)
    except Exception as e: s["log"].append({"t": now, "eur": 0, "why": f"Retire-Fehler: {e}"})
    for p in [x for x in s["products"] if not x.get("retired")]:
        try: ensure_payment_link(p)
        except Exception as e: s["log"].append({"t": now, "eur": 0, "why": f"Stripe-Fehler {p['slug']}: {e}"})
        if p.get("content_html") is not None or p.get("pay_url"): page(p)
    try: count_sales(s)
    except Exception as e: s["log"].append({"t": now, "eur": 0, "why": f"Verkaufszaehlung-Fehler: {e}"})
    try: freelancer_step(s)
    except Exception as e: s["log"].append({"t": now, "eur": 0, "why": f"Freelancer-Fehler: {e}"})
    if s["balance_eur"] <= 0: s["alive"] = False
    render(s); STATE_F.write_text(json.dumps(s, indent=2, ensure_ascii=False))
    print(f"Kontostand {s['balance_eur']} EUR, Produkte {len(s['products'])}")

if __name__ == "__main__": main()
