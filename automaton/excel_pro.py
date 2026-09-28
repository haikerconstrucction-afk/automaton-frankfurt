"""Haiktec Excel Pro: Startseite mit Logo & Navigation, formatierte Tabellen, Dropdowns, Ampel-Formatierung, Dashboard mit KPIs und Diagramm."""
import re, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, DataBarRule
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.utils import get_column_letter
ROOT = pathlib.Path(__file__).resolve().parent.parent
G, DG, LG, GR = "0B6E4F", "0F3D2E", "E7F3EE", "6B7A73"
FMT = {"eur": '#,##0.00 "€"', "usd": '"$"#,##0.00', "number": "#,##0.##", "int": "0", "percent": "0.0%", "date": "DD.MM.YYYY", "hours": "0.00"}
L = {"de": ("Start", "Übersicht", "So funktioniert's", "Kennzahlen", "Zum Blatt", "Nur weiße Zellen ausfüllen – grüne Spalten rechnen automatisch.", "Erstellt von Haiktec · haiktec.tech"),
     "en": ("Start", "Dashboard", "How it works", "Key figures", "Go to sheet", "Only fill in white cells – green columns calculate automatically.", "Made by Haiktec · haiktec.tech"),
     "es": ("Inicio", "Resumen", "Cómo funciona", "Indicadores", "Ir a la hoja", "Rellene solo las celdas blancas: las columnas verdes se calculan solas.", "Creado por Haiktec · haiktec.tech"),
     "fr": ("Accueil", "Tableau de bord", "Mode d'emploi", "Indicateurs", "Aller à la feuille", "Remplissez uniquement les cellules blanches – les colonnes vertes se calculent seules.", "Créé par Haiktec · haiktec.tech"),
     "zh": ("开始", "总览", "使用说明", "关键指标", "打开工作表", "只需填写白色单元格，绿色列会自动计算。", "由 Haiktec 制作 · haiktec.tech")}
thin = Side(style="thin", color="D5DDD9")

def safe_name(n, used):
    n = re.sub(r"[\[\]\*\?/\\:]", "-", str(n))[:28] or "Daten"
    b = n; i = 2
    while n in used: n = f"{b[:25]}-{i}"; i += 1
    used.add(n); return n

def conv(x, typ):
    if isinstance(x, str) and x.startswith("="): return x
    if typ in ("eur", "usd", "number", "int", "percent", "hours") and isinstance(x, str):
        t = x.replace("€", "").replace("$", "").replace("%", "").strip().replace(",", ".")
        try: v = float(t); return v / 100 if typ == "percent" and "%" in x else v
        except ValueError: return x
    if typ == "date" and isinstance(x, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", x):
        import datetime; return datetime.date.fromisoformat(x)
    return x

def build(p, path):
    lang = p.get("language", "de"); T = L.get(lang, L["en"])
    wb = Workbook(); st = wb.active; st.title = T[0]; used = {T[0], T[1]}
    st.sheet_view.showGridLines = False
    st.column_dimensions["A"].width = 3; st.column_dimensions["B"].width = 11; st.column_dimensions["C"].width = 95
    try:
        from openpyxl.drawing.image import Image as XImg
        im = XImg(str(ROOT/"site/icon.png")); im.width = im.height = 64; st.add_image(im, "B2")
    except Exception: pass
    st.row_dimensions[2].height = 50
    st["C2"] = p["title"]; st["C2"].font = Font(bold=True, size=18, color=DG); st["C2"].alignment = Alignment(wrap_text=True, vertical="center")
    st["C3"] = p.get("pitch", ""); st["C3"].font = Font(size=11, color=GR); st["C3"].alignment = Alignment(wrap_text=True)
    st.row_dimensions[3].height = 45
    r = 5; st.cell(r, 3, T[2]).font = Font(bold=True, size=13, color=G); r += 1
    for i, g in enumerate(p.get("guide", [])[:10], 1):
        st.cell(r, 2, f"{i}."); st.cell(r, 2).font = Font(bold=True, color=G); st.cell(r, 2).alignment = Alignment(horizontal="right", vertical="top")
        c = st.cell(r, 3, g); c.alignment = Alignment(wrap_text=True, vertical="top"); st.row_dimensions[r].height = max(18, 15 * (len(g) // 80 + 1)); r += 1
    r += 1; st.cell(r, 3, T[5]).font = Font(italic=True, color=GR); r += 2
    nav_row = r; r += 1
    formula_cols, sheets = {}, []
    for sh in p.get("sheets", [])[:6]:
        name = safe_name(sh.get("name", "Daten"), used); ws = wb.create_sheet(name); sheets.append((name, ws, sh))
        cols = [c for c in sh.get("columns", []) if str(c).strip()]
        if not cols: wb.remove(ws); sheets.pop(); continue
        types = sh.get("types") or ["text"] * len(cols)
        types = (types + ["text"] * len(cols))[:len(cols)]
        heads = [str(c)[:40] or f"Spalte{i+1}" for i, c in enumerate(cols)]
        seen = set(); heads = [h if not (h in seen or seen.add(h)) else f"{h} {i}" for i, h in enumerate(heads)]
        ws.append(heads)
        rows = sh.get("rows", [])[:300]
        blank = min(int(sh.get("blank_rows", 30) or 30), 200)
        fcols = {}
        tm = {}
        for cname, f in (sh.get("formulas") or {}).items():
            if cname in heads and isinstance(f, str) and f.startswith("="): tm[heads.index(cname)] = f
        if tm:
            rows = [list(rr) + [None] * (len(cols) - len(rr)) for rr in rows]
            for i, rr in enumerate(rows):
                for j, f in tm.items(): rr[j] = f.replace("{r}", str(i + 2))
        for rr in rows:
            for j, x in enumerate(rr[:len(cols)]):
                if isinstance(x, str) and x.startswith("="): fcols.setdefault(j, x)
        for rr in rows:
            ws.append([conv(x, types[j] if j < len(types) else "text") for j, x in enumerate(rr[:len(cols)])])
        # Formeln in Leerzeilen fortschreiben
        last_data = ws.max_row
        for k in range(blank):
            rownum = last_data + 1 + k
            for j, f in fcols.items():
                src_row = last_data
                tmpl = ws.cell(src_row, j + 1).value
                if isinstance(tmpl, str) and tmpl.startswith("="):
                    expr = re.sub(r"([A-Z]{1,2})" + str(src_row) + r"\b", lambda m: m.group(1) + str(rownum), tmpl)[1:]
                    ws.cell(rownum, j + 1, f'=IF(A{rownum}="","",{expr})')
            if not fcols: ws.cell(rownum, 1, None)
        end = max(ws.max_row, 2)
        ref = f"A1:{get_column_letter(len(cols))}{end}"
        tname = re.sub(r"[^A-Za-z0-9_]", "_", "T_" + name)[:30]
        t = Table(displayName=tname, ref=ref); t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium7", showRowStripes=True)
        ws.add_table(t)
        for j, h in enumerate(heads, 1):
            c = ws.cell(1, j); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=G)
            c.alignment = Alignment(wrap_text=True, vertical="center"); typ = types[j - 1]
            width = max(12, min(40, max(len(h), *(len(str(ws.cell(i, j).value or "")) for i in range(2, min(end, 30) + 1))) + 3))
            ws.column_dimensions[get_column_letter(j)].width = width
            for i in range(2, end + 1):
                cell = ws.cell(i, j); cell.border = Border(bottom=thin)
                if typ in FMT: cell.number_format = FMT[typ]
                if (j - 1) in fcols: cell.fill = PatternFill("solid", fgColor=LG); cell.font = Font(color=DG)
            if typ in ("eur", "usd", "number"):
                ws.conditional_formatting.add(f"{get_column_letter(j)}2:{get_column_letter(j)}{end}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C0392B", bold=True)))
            if typ == "percent":
                ws.conditional_formatting.add(f"{get_column_letter(j)}2:{get_column_letter(j)}{end}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="5FB38F"))
        for dv in sh.get("dropdowns", [])[:4]:
            try:
                j = heads.index(dv["column"]) + 1; opts = ",".join(str(o).replace(",", " ")[:30] for o in dv["options"][:12])
                if len(opts) < 250:
                    v = DataValidation(type="list", formula1=f'"{opts}"', allow_blank=True); ws.add_data_validation(v)
                    v.add(f"{get_column_letter(j)}2:{get_column_letter(j)}{end}")
            except Exception: pass
        ws.row_dimensions[1].height = 32; ws.freeze_panes = "A2"
        ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.center.text = T[6]
        sh["_end"], sh["_heads"], sh["_name"] = end, heads, name
    # Navigation
    st.cell(nav_row, 3, T[4]).font = Font(bold=True, size=13, color=G)
    for k, nm in enumerate([T[1]] + [n for n, _, _ in sheets]):
        c = st.cell(nav_row + 1 + k, 3, "→ " + nm); c.hyperlink = f"#'{nm}'!A1"; c.font = Font(color=G, underline="single")
    st.cell(nav_row + 3 + len(sheets), 3, T[6]).font = Font(size=9, color=GR)
    # Dashboard
    db = wb.create_sheet(T[1], 1); db.sheet_view.showGridLines = False
    db.column_dimensions["A"].width = 3; db.column_dimensions["B"].width = 36; db.column_dimensions["C"].width = 22
    db["B2"] = T[1] + " – " + p["title"][:60]; db["B2"].font = Font(bold=True, size=16, color=DG)
    db["B4"] = T[3]; db["B4"].font = Font(bold=True, size=12, color=G)
    names = {sh.get("name"): sh["_name"] for _, _, sh in sheets}
    def fix(f):
        for orig, new in names.items():
            if orig and orig != new: f = f.replace(f"'{orig}'!", f"'{new}'!").replace(f"{orig}!", f"'{new}'!")
        return f
    kr = 5
    for k in (p.get("kpis") or [])[:8]:
        if not isinstance(k, dict) or not str(k.get("formula", "")).startswith("="): continue
        a = db.cell(kr, 2, str(k.get("label", ""))[:50]); a.fill = PatternFill("solid", fgColor=LG); a.font = Font(bold=True, color=DG)
        b = db.cell(kr, 3, fix(k["formula"])); b.fill = PatternFill("solid", fgColor=LG); b.font = Font(bold=True, size=13, color=G)
        b.number_format = FMT.get(k.get("type", "number"), "#,##0.00"); b.alignment = Alignment(horizontal="right")
        a.border = b.border = Border(bottom=Side(style="medium", color="FFFFFF")); db.row_dimensions[kr].height = 24; kr += 1
    ch = p.get("chart") or {}
    try:
        tgt = next(s3 for s3 in sheets if s3[2].get("name") == ch.get("sheet")) if ch.get("sheet") else sheets[0]
        name, ws, sh = tgt; heads = sh["_heads"]
        cat = heads.index(ch["category_column"]) + 1 if ch.get("category_column") in heads else 1
        val = heads.index(ch["value_column"]) + 1 if ch.get("value_column") in heads else None
        if val is None:
            val = next((j + 1 for j, t in enumerate((sh.get("types") or [])) if t in ("eur", "usd", "number", "hours", "int")), None)
        if val:
            n = min(sh["_end"], len(sh.get("rows", [])) + 1)
            kind = ch.get("type", "bar"); c = PieChart() if kind == "pie" else LineChart() if kind == "line" else BarChart()
            c.title = ch.get("title") or heads[val - 1]; c.height, c.width = 8, 16
            c.add_data(Reference(ws, min_col=val, min_row=1, max_row=n), titles_from_data=True)
            c.set_categories(Reference(ws, min_col=cat, min_row=2, max_row=n))
            if kind != "pie": c.legend = None
            try: c.series[0].graphicalProperties.solidFill = G
            except Exception: pass
            db.add_chart(c, f"E4")
    except Exception as e: print("Diagramm:", e)
    for w in (st, db):
        w.page_setup.orientation = "landscape"; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 1; w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.active = 0
    wb.properties.creator = "Haiktec"; wb.properties.title = p["title"]
    wb.save(path)
    f = sum(1 for ws in wb.worksheets for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("="))
    return f
